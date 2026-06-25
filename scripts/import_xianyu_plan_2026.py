"""将「湖南-2026招生计划版（不含艺术类）.xlsx」导入 SQLite。

源表第 0 行是分组标题，第 1 行是真实列名，数据从第 2 行开始。
"""

import sqlite3
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "raw_data" / "闲鱼数据" / "湖南-2026招生计划版（不含艺术类）.xlsx"
DB = ROOT / "output" / "hunan_gkdata.db"
TABLE = "招生计划2026"


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = ws.iter_rows(values_only=True)

    next(rows)  # 跳过分组标题行
    header = list(next(rows))  # 真实列名

    # 处理重复/空列名
    seen: dict[str, int] = {}
    columns: list[str] = []
    for i, name in enumerate(header):
        col = str(name).strip() if name is not None else f"col_{i}"
        if col in seen:
            seen[col] += 1
            col = f"{col}_{seen[col]}"
        else:
            seen[col] = 0
        columns.append(col)

    conn = sqlite3.connect(DB)
    cur = conn.cursor()
    cur.execute(f'DROP TABLE IF EXISTS "{TABLE}"')
    cols_def = ", ".join(f'"{c}"' for c in columns)
    cur.execute(f'CREATE TABLE "{TABLE}" ({cols_def})')

    placeholders = ", ".join("?" for _ in columns)
    insert_sql = f'INSERT INTO "{TABLE}" VALUES ({placeholders})'

    n = len(columns)
    batch = []
    total = 0
    for row in rows:
        row = list(row)
        if len(row) < n:
            row += [None] * (n - len(row))
        elif len(row) > n:
            row = row[:n]
        # 全空行跳过
        if all(v is None for v in row):
            continue
        batch.append(row)
        if len(batch) >= 1000:
            cur.executemany(insert_sql, batch)
            total += len(batch)
            batch.clear()
    if batch:
        cur.executemany(insert_sql, batch)
        total += len(batch)

    conn.commit()

    # 加常用索引
    for idx_col in ("院校名称", "院校代码", "专业组代码", "批次", "科类"):
        if idx_col in columns:
            cur.execute(
                f'CREATE INDEX IF NOT EXISTS "idx_{idx_col}" ON "{TABLE}" ("{idx_col}")'
            )
    conn.commit()
    conn.close()

    print(f"导入完成：{total} 行 -> {DB} (表 {TABLE})")


if __name__ == "__main__":
    main()
