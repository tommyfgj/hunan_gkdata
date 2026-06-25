"""将官方艺术类平行组投档线(2024/2025)整合并写入 output/hunan_gkdata.db。

源文件(湖南考试院官方)：
- raw_data/湖南考试院网站官方数据/2024年艺术类平行组投档线.xlsx (15列, 无"再选次高")
- raw_data/湖南考试院网站官方数据/2025年艺术类平行组投档线.xlsx (16列)

两份文件前 2 行为标题/说明，第 3 行为列名，数据从第 4 行起。
统一为一张表 艺术类平行组投档线，含 年份 列。
"""

import sqlite3
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "raw_data" / "湖南考试院网站官方数据"
DB = ROOT / "output" / "hunan_gkdata.db"
TABLE = "艺术类平行组投档线"

COLUMNS = [
    "年份", "科类", "计划类别", "院校代号", "院校名称", "专业组编号", "专业组名称",
    "投档线", "文化成绩", "语数之和", "语数最高", "外语", "首选科目",
    "再选最高", "再选次高", "志愿序号", "备注",
]

SOURCES = {
    2024: "2024年艺术类平行组投档线.xlsx",
    2025: "2025年艺术类平行组投档线.xlsx",
}


def read_year(year: int, filename: str) -> list[tuple]:
    wb = openpyxl.load_workbook(RAW / filename, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[2]]
    idx = {name: i for i, name in enumerate(header)}

    records: list[tuple] = []
    for row in rows[3:]:
        if not row or row[idx["院校代号"]] in (None, ""):
            continue

        def get(col: str):
            i = idx.get(col)
            return row[i] if i is not None else None

        records.append((
            year, get("科类"), get("计划类别"), get("院校代号"), get("院校名称"),
            get("专业组编号"), get("专业组名称"), get("投档线"), get("文化成绩"),
            get("语数之和"), get("语数最高"), get("外语"), get("首选科目"),
            get("再选最高"), get("再选次高"), get("志愿序号"), get("备注"),
        ))
    return records


def main() -> None:
    conn = sqlite3.connect(DB)
    cur = conn.cursor()
    cur.execute(f'DROP TABLE IF EXISTS "{TABLE}"')
    cols_def = ", ".join(f'"{c}"' for c in COLUMNS)
    cur.execute(f'CREATE TABLE "{TABLE}" ({cols_def})')

    placeholders = ", ".join("?" for _ in COLUMNS)
    insert_sql = f'INSERT INTO "{TABLE}" VALUES ({placeholders})'

    total = 0
    for year, filename in SOURCES.items():
        records = read_year(year, filename)
        cur.executemany(insert_sql, records)
        total += len(records)
        print(f"  {year}: {len(records)} 行 (来源 {filename})")

    for name, cols in [
        ("idx_art_td_年份", "年份"),
        ("idx_art_td_院校代号", "院校代号"),
        ("idx_art_td_专业组名称", "专业组名称"),
        ("idx_art_td_年份专业组", "年份, 专业组名称"),
    ]:
        cur.execute(f'CREATE INDEX IF NOT EXISTS "{name}" ON "{TABLE}" ({cols})')

    conn.commit()

    print(f"\n=== 校验 ===")
    cur.execute(f'SELECT 年份, COUNT(*) FROM "{TABLE}" GROUP BY 年份 ORDER BY 年份')
    print("  年份分布:", dict(cur.fetchall()))
    cur.execute(
        f"SELECT 年份, COUNT(*) FROM \"{TABLE}\" WHERE 专业组名称 LIKE '%美术%' GROUP BY 年份 ORDER BY 年份"
    )
    print("  美术与设计类:", dict(cur.fetchall()))
    conn.close()
    print(f"\n共写入 {total} 行 -> {DB}")


if __name__ == "__main__":
    main()
